import tempfile
import unittest
from pathlib import Path

from sqlalchemy.ext.asyncio import create_async_engine, async_sessionmaker
from models import Base, User, Submission, Evaluation, SubmissionStatus
from utils import calculate_average_scores, update_submission_scores, export_to_excel


class ScoringTests(unittest.IsolatedAsyncioTestCase):
    async def asyncSetUp(self):
        self.engine = create_async_engine("sqlite+aiosqlite:///:memory:")
        async with self.engine.begin() as conn:
            await conn.run_sync(Base.metadata.create_all)
        self.session = async_sessionmaker(self.engine, expire_on_commit=False)()
        self.user = User(telegram_id=1, nickname="participant", points=0, evaluated_submissions_count=0)
        self.admin = User(telegram_id=2, nickname="judge", is_admin=True)
        self.submission = Submission(user=self.user, photo_file_id="photo")
        self.session.add_all([self.user, self.admin, self.submission])
        await self.session.flush()

    async def asyncTearDown(self):
        await self.session.close()
        await self.engine.dispose()

    async def evaluation(self, visual=0, creativity=5, technique=5, harmony=5):
        result = Evaluation(submission_id=self.submission.id, admin_id=self.admin.id,
                            visual_quality=visual, creativity=creativity,
                            technical_skills=technique, overall_harmony=harmony,
                            total_score=visual+creativity+technique+harmony)
        self.session.add(result)
        await self.session.flush()
        return result

    async def test_zero_visual_does_not_discard_other_scores(self):
        await self.evaluation()
        result = await calculate_average_scores(self.session, self.submission.id)
        self.assertEqual(result["visual_quality"], 0)
        self.assertEqual(result["total_score"], 15)

    async def test_recalculation_does_not_double_credit_points(self):
        evaluation = await self.evaluation(visual=5)
        await update_submission_scores(self.session, self.submission.id)
        await update_submission_scores(self.session, self.submission.id)
        self.assertEqual(self.user.points, 20)
        self.assertEqual(self.user.evaluated_submissions_count, 1)
        evaluation.total_score = 24
        await self.session.flush()
        await update_submission_scores(self.session, self.submission.id)
        self.assertEqual(self.user.points, 24)
        self.assertEqual(self.user.evaluated_submissions_count, 1)

    async def test_rejecting_accepted_result_reverses_credit_once(self):
        from utils import apply_submission_result
        await self.evaluation(visual=5)
        await update_submission_scores(self.session, self.submission.id)
        apply_submission_result(self.submission, {}, accepted=False)
        apply_submission_result(self.submission, {}, accepted=False)
        self.assertEqual(self.user.points, 0)
        self.assertEqual(self.user.evaluated_submissions_count, 0)
        self.assertEqual(self.submission.status, SubmissionStatus.REJECTED)

    async def test_zero_score_is_a_valid_evaluated_result(self):
        await self.evaluation(0, 0, 0, 0)
        await update_submission_scores(self.session, self.submission.id)
        self.assertEqual(self.submission.status, SubmissionStatus.EVALUATED)
        self.assertEqual(self.user.points, 0)
        self.assertEqual(self.user.evaluated_submissions_count, 1)

    async def test_missing_evaluations_leave_submission_pending(self):
        await update_submission_scores(self.session, self.submission.id)
        self.assertEqual(self.submission.status, SubmissionStatus.PENDING)
        self.assertEqual(self.user.evaluated_submissions_count, 0)

    async def test_export_preserves_fractional_scores(self):
        from openpyxl import load_workbook
        await self.evaluation(visual=5.5)
        await update_submission_scores(self.session, self.submission.id)
        with tempfile.TemporaryDirectory() as temp:
            filename = str(Path(temp) / "results.xlsx")
            await export_to_excel(self.session, filename)
            book = load_workbook(filename)
            try:
                for sheet, column in [("Работы", "Визуал"), ("Оценки", "Visual")]:
                    rows = list(book[sheet].values)
                    self.assertEqual(rows[1][rows[0].index(column)], 5.5)
            finally:
                book.close()

    async def test_export_keeps_formula_like_nicknames_as_text(self):
        from openpyxl import load_workbook
        self.user.nickname = "=1+1"
        await self.session.flush()
        with tempfile.TemporaryDirectory() as temp:
            filename = str(Path(temp) / "results.xlsx")
            await export_to_excel(self.session, filename)
            book = load_workbook(filename)
            try:
                sheet = book["Участники"]
                headers = [cell.value for cell in sheet[1]]
                cell = sheet.cell(2, headers.index("Ник") + 1)
                self.assertEqual(cell.value, "=1+1")
                self.assertEqual(cell.data_type, "s")
            finally:
                book.close()

    async def test_admin_handler_rejudges_without_double_credit(self):
        from types import SimpleNamespace
        from unittest.mock import AsyncMock, patch
        from routers.admin import judge_decide
        cb = SimpleNamespace(data="a:decide:accept", from_user=SimpleNamespace(id=2),
                             message=AsyncMock(), answer=AsyncMock())
        state = AsyncMock()
        data = {"submission_id": self.submission.id, "visual": 5, "creativity": 5,
                "technique": 5, "harmony": 5, "total": 20}
        state.get_data.return_value = data
        with patch("routers.admin.require_admin", new=AsyncMock(return_value=True)):
            await judge_decide(cb, self.session, state)
            data.update(visual=9, total=24)
            await judge_decide(cb, self.session, state)
        self.assertEqual(self.user.points, 24)
        self.assertEqual(self.user.evaluated_submissions_count, 1)
        self.assertEqual(self.submission.total_score, 24)

    async def test_invalid_criteria_do_not_mutate_submission_or_points(self):
        from utils import apply_submission_result
        scores = dict(visual_quality=5, creativity=5, technical_skills=5, overall_harmony=5, total_score=20)
        for invalid in (-1, 11, float("nan"), float("inf"), None):
            with self.subTest(value=invalid):
                with self.assertRaises(ValueError):
                    apply_submission_result(self.submission, dict(scores, visual_quality=invalid))
                self.assertEqual(self.user.points, 0)
                self.assertEqual(self.submission.status, SubmissionStatus.PENDING)

    async def test_submission_window_rejects_nonpositive_and_fractional_weeks(self):
        from utils import check_submission_limit
        for weeks in (0, -1, 1.5, True):
            with self.subTest(weeks=weeks), self.assertRaises(ValueError):
                await check_submission_limit(self.session, self.user.id, weeks)
        self.assertEqual(await check_submission_limit(self.session, self.user.id), (True, 1))

    async def test_equal_points_have_stable_ordinal_ranks(self):
        from utils import get_user_ranking
        self.assertEqual(await get_user_ranking(self.session, self.user.id), 1)
        self.assertEqual(await get_user_ranking(self.session, self.admin.id), 2)

    async def test_banned_users_do_not_occupy_leaderboard_places(self):
        from utils import get_user_ranking
        from routers.user import leaderboard_cmd
        from unittest.mock import AsyncMock
        self.admin.is_banned = True
        self.admin.points = 999
        await self.session.flush()
        self.assertEqual(await get_user_ranking(self.session, self.user.id), 1)
        self.assertEqual(await get_user_ranking(self.session, self.admin.id), 0)
        message = AsyncMock()
        await leaderboard_cmd(message, self.session)
        self.assertNotIn("judge", message.answer.call_args.args[0])
